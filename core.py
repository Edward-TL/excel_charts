```python
"""core.py

Base definitions for the chart library.
"""

from __future__ import annotations

import abc
from dataclasses import dataclass, field
from enum import Enum
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.chart import Chart


class MoneyAxis(str, Enum):
    """Enum to indicate which axis represents monetary values."""

    X = "x"
    Y = "y"


@dataclass
class ColorPalette:
    """Simple color palette definition for charts.

    Attributes
    ----------
    primary: str
        Primary color in hex (e.g., "#4A90E2").
    secondary: str
        Secondary color in hex.
    accent: str
        Accent color in hex.
    """

    primary: str = "#4A90E2"
    secondary: str = "#50E3C2"
    accent: str = "#F5A623"


class BaseChart(abc.ABC):
    """Abstract base class for all chart types.

    All concrete chart classes inherit from this to gain common attributes and
    helper methods.
    """

    def __init__(
        self,
        data: pd.DataFrame,
        sheet: str = "Sheet1",
        position: str = "A1",
        color_palette: ColorPalette | None = None,
        money_axis: MoneyAxis = MoneyAxis.Y,
    ) -> None:
        self.data = data
        self.sheet = sheet
        self.position = position
        self.color_palette = color_palette or ColorPalette()
        self.money_axis = money_axis

    @property
    def _workbook(self) -> Workbook:
        """Placeholder property – the workbook must be supplied when adding the chart.
        Sub‑classes typically receive a ``Workbook`` instance via ``add_to_workbook``.
        """
        raise NotImplementedError

    @property
    def _worksheet(self) -> Worksheet:
        """Return the worksheet object where the chart will be placed."""
        raise NotImplementedError

    @abc.abstractmethod
    def _create_chart(self) -> Chart:
        """Create and configure the specific OpenPyXL chart instance.

        Returns
        -------
        Chart
            An instantiated OpenPyXL chart ready to be added to a worksheet.
        """
        pass

    def add_to_workbook(self, wb: Workbook) -> Chart:
        """Add the chart to the supplied workbook.

        Parameters
        ----------
        wb : Workbook
            The OpenPyXL workbook where the chart should be inserted.

        Returns
        -------
        Chart
            The chart object that was added to the worksheet.
        """
        # Resolve worksheet lazily – subclasses may need the workbook to locate sheet.
        self._wb = wb
        if self.sheet not in wb.sheetnames:
            wb.create_sheet(self.sheet)
        self._ws = wb[self.sheet]
        chart = self._create_chart()
        self._ws.add_chart(chart, self.position)
        return chart
```
